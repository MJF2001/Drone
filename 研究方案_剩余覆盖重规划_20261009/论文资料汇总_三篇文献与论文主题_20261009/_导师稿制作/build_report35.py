from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parent.parent
STEM='35_导师汇报稿_实验设计图解版_20261011'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.72)
sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)

def font(style,size,east='宋体',bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size)
    style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    rf.set(qn('w:eastAsia'),east)
    for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rf.attrib.pop(qn('w:'+attr),None)

font(doc.styles['Normal'],12)
pf=doc.styles['Normal'].paragraph_format
pf.line_spacing=1.3;pf.space_after=Pt(5);pf.first_line_indent=Pt(24)
pf.widow_control=True
font(doc.styles['Title'],18,'黑体',True)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].paragraph_format.first_line_indent=Pt(0)
for n,size in [('Heading 1',14),('Heading 2',12)]:
    font(doc.styles[n],size,'黑体',True)
    f=doc.styles[n].paragraph_format;f.space_before=Pt(9);f.space_after=Pt(6)
    f.first_line_indent=Pt(0);f.keep_with_next=True;f.keep_together=True
font(doc.styles['Header'],9,'宋体');font(doc.styles['Footer'],9,'宋体')
hp=sec.header.paragraphs[0];hp.text='无人机剩余覆盖路径重规划研究进展'
hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.paragraph_format.first_line_indent=Pt(0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent=Pt(0)
fp.add_run('第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
fp.add_run(' 页')

def table(lines):
    data=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    data=[r for r in data if not all(set(v)<=set('-: ') for v in r)]
    count=len(data[0]);widths={2:[1.35,5.45],3:[1.35,3.95,1.5],4:[2.3,1.5,1.5,1.5],5:[2.2,1.15,1.15,1.15,1.15]}[count]
    is_rationale=data[0][-1]=='设置目的'
    if is_rationale:widths=[1.25,2.8,2.75]
    if data[0][0]=='实验数据':widths=[1.5,.65,4.65]
    t=doc.add_table(rows=0,cols=count);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(data):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,value in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF2');tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if is_rationale or count==2 or (count==3 and j<2) or (data[0][0]=='实验数据' and j==2) or (count!=3 and j==0) else WD_ALIGN_PARAGRAPH.CENTER
            if (is_rationale or data[0][0]=='预算') and ri < len(data)-1:p.paragraph_format.keep_with_next=True
            r=p.add_run(text);r.font.size=Pt(10.5);r.bold=ri==0
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(2);after.paragraph_format.line_spacing=1
    after.paragraph_format.first_line_indent=Pt(0);after.add_run().font.size=Pt(2)

lines=(ROOT/(STEM+'.md')).read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line=='<!-- pagebreak -->':doc.add_page_break();continue
    if line.startswith('!['):
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(Path(__file__).resolve().parent/line.rsplit('(',1)[1].rstrip(')')),width=Inches(6.8))
        continue
    if line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
        table(block);continue
    if line.startswith('# '):
        title=line[2:].replace('条件下无人机','条件下\n无人机')
        p=doc.add_paragraph(title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line=='2026年10月11日':
        p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(12)
        p.runs[0].font.size=Pt(10.5)
    else:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        import re
        for fragment in re.split(r'(\*\*.*?\*\*)',line):
            r=p.add_run(fragment[2:-2] if fragment.startswith('**') else fragment);r.bold=fragment.startswith('**')
        if line.startswith('优化前后在同批') or line.startswith('24个新快照、相同5秒预算下'):p.paragraph_format.keep_with_next=True
        if re.match(r'图[0-9]+ ',line):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
            for r in p.runs:r.font.size=Pt(10.5)
        if line.startswith('**'):
            p.paragraph_format.first_line_indent=Pt(0)
        if line.startswith('**第') and line.endswith('**'):
            p.paragraph_format.keep_with_next=True
        if line.startswith('['):
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.1
            for r in p.runs:r.font.size=Pt(10.5)
for parent in (doc.styles.element,doc.element):
    for border in list(parent.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='施工待检区域变化条件下无人机剩余覆盖路径重规划研究进展'
doc.core_properties.subject='研究范围 文献依据 对照算法与阶段性实验'
doc.core_properties.author=''
doc.save(ROOT/(STEM+'.docx'))
print(ROOT/(STEM+'.docx'))



